from __future__ import annotations
from datetime import datetime
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from core.paths import OUTPUT_DIR
from agent.risk_brief import baseline_brief, scenario_brief

CJK='Noto Sans CJK SC'

def _clean_md(s: str) -> str:
    s=re.sub(r'\*\*(.*?)\*\*',r'\1',s)
    s=s.replace('`','')
    return s.strip()

def _add_brief(doc: Document, text: str):
    for raw in (text or '').splitlines():
        line=raw.strip()
        if not line: continue
        if line.startswith('### '):
            doc.add_heading(_clean_md(line[4:]), level=2)
        elif line.startswith('- '):
            doc.add_paragraph(_clean_md(line[2:]), style='List Bullet')
        else:
            doc.add_paragraph(_clean_md(line))

def _force_cjk(doc: Document):
    for st in doc.styles:
        try:
            st.font.name=CJK
            st._element.rPr.rFonts.set(qn('w:eastAsia'),CJK)
        except Exception: pass
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.name=CJK
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),CJK)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.name=CJK
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),CJK)

def export_risk_report(baseline_result: dict, scenario_result: dict|None, validation: dict|None, snapshot_id: str='') -> str:
    doc=Document()
    doc.styles['Normal'].font.name=CJK; doc.styles['Normal'].font.size=Pt(10.5)
    doc.add_heading('Water Risk Copilot 风险管理简报',0)
    doc.add_paragraph(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    if snapshot_id: doc.add_paragraph(f'Input Snapshot：{snapshot_id}')
    doc.add_heading('1. 数据质量与计算状态',1)
    if validation:
        doc.add_paragraph(f"状态：{validation.get('status')}")
        doc.add_paragraph('Coverage：'+str(validation.get('coverage')))
        for k,cn in [('issues','冲突/阻断'),('warnings','警告'),('data_gaps','数据缺口')]:
            vals=validation.get(k,[])
            if vals: doc.add_paragraph(cn+'：'+'；'.join(map(str,vals)))
    doc.add_heading('2. Baseline 结果与解释',1)
    _add_brief(doc, baseline_brief(baseline_result))
    if baseline_result and baseline_result.get('data') is not None:
        s=baseline_result['data']['summary']
        table=doc.add_table(rows=1,cols=6); table.style='Table Grid'
        hdr=table.rows[0].cells
        for i,x in enumerate(['材料','PRWI','Coverage','下界','上界','Unknown']): hdr[i].text=x
        for _,r in s.iterrows():
            c=table.add_row().cells
            vals=[r.material,f'{r.PRWI:.4f}',f'{r.Coverage*100:.1f}%',f'{r.PRWI_lower:.4f}',f'{r.PRWI_upper:.4f}',f'{r.unknown_share_U:.4f}']
            for i,v in enumerate(vals): c[i].text=str(v)
    doc.add_heading('3. Scenario 结果与解释',1)
    _add_brief(doc, scenario_brief(scenario_result) if scenario_result else '本次未运行 Scenario。')
    doc.add_heading('4. 管理建议与边界',1)
    doc.add_paragraph('优先核验高贡献节点的真实采购信息；对高风险节点进行持续监测；必要时配置库存与替代供应。所有真实采购调整、供应商替换和投资决策均需人工确认。')
    doc.add_paragraph('PRWI 为相对风险筛查指数，不是损失概率或财务预测。Future demo、OA、库存/替代等情景假设必须显式披露。')
    doc.add_heading('5. 版本与可追溯信息',1)
    doc.add_paragraph('Baseline model: v2.0；Scenario engine: v2.0；Data bundle: 2026-09-06-integrated-mvp。')
    _force_cjk(doc)
    out=OUTPUT_DIR/f'Water_Risk_Brief_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
    doc.save(out)
    return str(out)
