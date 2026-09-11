from __future__ import annotations
import io, re, json, difflib
from pathlib import Path
from typing import Dict, List, Tuple
import pandas as pd
from docx import Document
from pypdf import PdfReader

from core.paths import ENGINE_DIR
import importlib.util

spec=importlib.util.spec_from_file_location('baseline_v2', ENGINE_DIR/'baseline_water_risk_v2.py')
b=importlib.util.module_from_spec(spec); spec.loader.exec_module(b)

CANONICAL=['enterprise','material','node_id','node_name','purchase_weight','year','source_type','confidence']
ALIASES={
 'enterprise':['enterprise','企业','公司','company'],
 'material':['material','原材料','材料','raw material','commodity'],
 'node_id':['node_id','节点编码','节点id','id'],
 'node_name':['node_name','供应节点','节点','产区','地区','location','supplier location','source country','country'],
 'purchase_weight':['purchase_weight','w','采购权重','采购占比','采购比例','purchase %','purchase_share','share','weight'],
 'year':['year','年份','参考期','period'],
}

def known_nodes_df():
    rows=[]
    for nid,h in b.HAZARD_DATA.items():
        mat='甘蔗' if nid.startswith('SC') else ('甜菜' if nid.startswith('SB') else '大豆')
        rows.append({'node_id':nid,'node_name':h['name'],'material':mat,'data_type':h['data_type'],'hazard_confidence':h['confidence'],'pfaf_id':h['pfaf_id']})
    return pd.DataFrame(rows)


def load_demo_procurement(material='甘蔗') -> pd.DataFrame:
    rows=[]
    for nid,e in b.EXPOSURE_DATA[material].items():
        if e['W']<=0: continue
        rows.append({'enterprise':b.ENTERPRISE_MAP[material],'material':material,'node_id':nid,
                     'node_name':b.HAZARD_DATA[nid]['name'],'purchase_weight':e['W'],'year':2025,
                     'source_type':'A','confidence':'Low'})
    return pd.DataFrame(rows)


def _extract_docx(path: str) -> str:
    d=Document(path)
    parts=[p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(parts)


def _extract_pdf(path: str) -> str:
    reader=PdfReader(path)
    return '\n'.join((p.extract_text() or '') for p in reader.pages)


def extract_candidates_from_text(text: str) -> pd.DataFrame:
    rows=[]
    mats=[m for m in ['甘蔗','甜菜','大豆'] if m in text]
    default_mat=mats[0] if len(mats)==1 else ''
    for nid,h in b.HAZARD_DATA.items():
        names=[h['name'], h['name'].split('(')[0], h['name'].split('-')[0]]
        pos=-1; matched=''
        for n in names:
            if n and n in text:
                pos=text.find(n); matched=n; break
        if pos<0: continue
        context=text[max(0,pos-80):pos+len(matched)+100]
        pct=re.search(r'(\d+(?:\.\d+)?)\s*%', context)
        w=float(pct.group(1))/100 if pct else None
        mat='甘蔗' if nid.startswith('SC') else ('甜菜' if nid.startswith('SB') else '大豆')
        rows.append({'enterprise':'','material':default_mat or mat,'node_id':nid,'node_name':h['name'],
                     'purchase_weight':w,'year':'','source_type':'V-candidate','confidence':'Low',
                     'evidence':context.replace('\n',' ')[:180]})
    return pd.DataFrame(rows)


def read_user_file(path: str):
    ext=Path(path).suffix.lower()
    if ext=='.csv':
        for enc in ['utf-8-sig','utf-8','gb18030']:
            try: return {'kind':'table','data':pd.read_csv(path,encoding=enc),'text':''}
            except Exception: pass
        raise ValueError('CSV 编码无法识别')
    if ext in ['.xlsx','.xls']:
        return {'kind':'table','data':pd.read_excel(path),'text':''}
    if ext=='.docx':
        text=_extract_docx(path); return {'kind':'text','data':extract_candidates_from_text(text),'text':text}
    if ext=='.pdf':
        text=_extract_pdf(path); return {'kind':'text','data':extract_candidates_from_text(text),'text':text}
    raise ValueError('仅支持 PDF、DOCX、XLSX、CSV')


def suggest_mapping(columns: List[str]) -> pd.DataFrame:
    cols=list(map(str,columns))
    rows=[]
    for target in ['enterprise','material','node_id','node_name','purchase_weight','year']:
        best=''
        lower={c:c.strip().lower() for c in cols}
        for c in cols:
            if lower[c] in [a.lower() for a in ALIASES[target]]:
                best=c; break
        if not best:
            for c in cols:
                if any(a.lower() in lower[c] or lower[c] in a.lower() for a in ALIASES[target]):
                    best=c; break
        rows.append({'system_field':target,'source_column':best,'required':target in ['material','purchase_weight']})
    return pd.DataFrame(rows)


def _best_node_match(name: str, material: str='') -> Tuple[str,str,float]:
    if not name: return '','',0
    k=known_nodes_df()
    if material in ['甘蔗','甜菜','大豆']:
        k=k[k.material==material]
    name=str(name).strip()
    # exact id/name first
    exact=k[(k.node_id.astype(str)==name)|(k.node_name.astype(str)==name)]
    if not exact.empty:
        r=exact.iloc[0]; return r.node_id,r.node_name,1.0
    scores=[]
    for _,r in k.iterrows():
        base=r.node_name.split('(')[0]
        score=max(difflib.SequenceMatcher(None,name,r.node_name).ratio(),difflib.SequenceMatcher(None,name,base).ratio())
        if name in r.node_name or base in name: score=max(score,0.9)
        scores.append((score,r.node_id,r.node_name))
    score,nid,nm=max(scores,default=(0,'',''))
    return (nid,nm,score) if score>=0.55 else ('','',score)


def apply_mapping(raw: pd.DataFrame, mapping: pd.DataFrame, weight_mode='0-1') -> pd.DataFrame:
    if raw is None or len(raw)==0: return pd.DataFrame(columns=CANONICAL)
    mapping=dict(zip(mapping['system_field'],mapping['source_column'])) if mapping is not None and len(mapping) else {}
    out=pd.DataFrame(index=raw.index)
    for target in ['enterprise','material','node_id','node_name','purchase_weight','year']:
        src=mapping.get(target,'')
        out[target]=raw[src] if src and src in raw.columns else ''
    out['purchase_weight']=pd.to_numeric(out['purchase_weight'],errors='coerce')
    if weight_mode=='百分数(0-100)': out['purchase_weight']=out['purchase_weight']/100.0
    out['source_type']='V-user-confirmed'
    out['confidence']='High'
    # fill node match if needed
    for i,row in out.iterrows():
        nid=str(row.get('node_id','') or '').strip()
        nname=str(row.get('node_name','') or '').strip()
        mat=str(row.get('material','') or '').strip()
        if nid not in b.HAZARD_DATA:
            mid,mname,score=_best_node_match(nname or nid,mat)
            if mid:
                out.at[i,'node_id']=mid
                out.at[i,'node_name']=mname
                if score<0.9: out.at[i,'confidence']='Medium'
    return out[CANONICAL]


def validate_normalized(df: pd.DataFrame) -> dict:
    issues=[]; warnings=[]; gaps=[]
    if df is None or len(df)==0:
        return {'status':'insufficient','issues':['没有可计算的采购记录'],'warnings':[],'data_gaps':['采购数据为空'],'coverage':0}
    if df['material'].replace('',pd.NA).isna().any(): issues.append('存在缺失 material')
    if df['purchase_weight'].isna().any(): issues.append('存在缺失/非数值 purchase_weight')
    invalid_nodes=[str(x) for x in df.loc[~df['node_id'].isin(b.HAZARD_DATA.keys()),'node_name'].tolist()]
    if invalid_nodes:
        gaps.append('以下节点无法匹配到当前16节点水风险库：'+', '.join(invalid_nodes))
    coverage_by_mat={}
    for mat,g in df.groupby('material'):
        s=float(pd.to_numeric(g['purchase_weight'],errors='coerce').fillna(0).sum())
        coverage_by_mat[str(mat)]=round(min(max(s,0),1),4)
        if s>1.0001: issues.append(f'{mat} 采购权重和={s:.4f}>1，请检查单位/重复记录')
        if s<0.9999: warnings.append(f'{mat} 已知采购权重和={s:.4f}，Unknown={max(0,1-s):.4f} 将保留，不自动归一化')
    if issues: status='conflict'
    elif invalid_nodes: status='partial'
    elif warnings: status='partial'
    else: status='success'
    return {'status':status,'issues':issues,'warnings':warnings,'data_gaps':gaps,'coverage':coverage_by_mat}
